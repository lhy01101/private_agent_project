"""
test_real_docs.py —— 用真实 python-docx / python-pptx 生成文件，验证 loader 切块

仅依赖已安装的库（python-docx, python-pptx），不依赖 chromadb/tree-sitter。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from file_ingest.loader import get_loader, DOCXLoader, PPTXLoader


def main():
    results = []
    def check(name, cond, detail=""):
        results.append(bool(cond))
        print(f"  {'✓' if cond else '✗'} {name}" + (f"  ({detail})" if detail else ""))

    base = Path("/tmp/file_ingest_real")
    base.mkdir(parents=True, exist_ok=True)

    # ---------- 生成真实 .docx ----------
    print("\n[1] DOCXLoader（真实 python-docx）")
    from docx import Document
    doc = Document()
    doc.add_heading("项目架构设计", level=1)
    doc.add_paragraph("本项目采用 RAG 架构，结合语义路由与 Chroma 向量库。")
    doc.add_heading("数据流程", level=2)
    doc.add_paragraph("用户上传文件后，系统解析并切块，小文件全文缓存，大文件写入向量库。")
    doc.add_heading("部署方案", level=2)
    doc.add_paragraph("部署在 Kubernetes 集群，使用 bge-m3 作为 embedding 模型。")
    docx_path = base / "design.docx"
    doc.save(str(docx_path))

    chunks = DOCXLoader().load(str(docx_path))
    check("非空", bool(chunks))
    check("按标题分层（多块）", len(chunks) >= 3)
    headings = [c.metadata.get("heading", "") for c in chunks]
    check("含标题信息", any(headings))
    check("含正文内容", any("RAG" in c.content or "向量" in c.content for c in chunks))
    print("    chunks 示例:")
    for c in chunks[:3]:
        print(f"      [{c.metadata.get('heading') or 'p'}] {c.content[:40]}")

    # ---------- 生成真实 .pptx ----------
    print("\n[2] PPTXLoader（真实 python-pptx）")
    from pptx import Presentation
    prs = Presentation()
    for i, (title, body) in enumerate([
        ("封面：季度总结", "2026 Q3 产品进展汇报"),
        ("用户增长", "DAU 提升 35%，新增文件管理模块"),
        ("技术亮点", "引入 Chroma 向量检索，支持 PDF/DOCX/PPTX"),
    ], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    pptx_path = base / "report.pptx"
    prs.save(str(pptx_path))

    chunks = PPTXLoader().load(str(pptx_path))
    check("非空", bool(chunks))
    check("每页一个 chunk", len(chunks) == 3, f"实际 {len(chunks)}")
    check("page metadata 正确", all(c.metadata.get("page") in (1, 2, 3) for c in chunks))
    check("内容完整", any("Chroma" in c.content for c in chunks))
    print("    chunks 示例:")
    for c in chunks:
        print(f"      [page {c.metadata.get('page')}] {c.content[:50]}")

    # ---------- get_loader 分发（文档类） ----------
    print("\n[3] get_loader 分发（文档类）")
    check("docx -> DOCXLoader", isinstance(get_loader(str(docx_path)), DOCXLoader))
    check("pptx -> PPTXLoader", isinstance(get_loader(str(pptx_path)), PPTXLoader))

    # ---------- ingest_file 对真实 docx 走 inline ----------
    print("\n[4] ingest_file 处理真实 .docx（inline 分支）")
    from file_ingest.tools import configure, ingest_file, query_file, _INLINE_CACHE
    from file_ingest.store import FileStore

    class _Coll:
        def __init__(self): self.docs = {}
        def count(self): return len(self.docs)
        def add(self, ids, embeddings, documents, metadatas):
            for i, id_ in enumerate(ids):
                self.docs[id_] = {"document": documents[i], "metadata": metadatas[i]}
        def query(self, qe, where, n_results=5):
            fid = where.get("file_id")
            hit = [d for d in self.docs.values() if d["metadata"].get("file_id") == fid]
            if not hit:
                return {"documents": [[]], "metadatas": [[]], "distances": [[]]}
            return {"documents": [[h["document"] for h in hit[:n_results]]],
                    "metadatas": [[h["metadata"] for h in hit[:n_results]]],
                    "distances": [[0.5] * len(hit[:n_results])]}
        def delete(self, where=None, ids=None): pass
        def get(self, where=None, limit=None):
            return {"ids": [], "documents": [], "metadatas": []}

    class _Client:
        def get_or_create_collection(self, name, metadata=None):
            if not hasattr(self, "c"):
                self.c = _Coll()
            return self.c

    def embed_fn(texts):
        import hashlib
        out = []
        for t in texts:
            h = hashlib.sha256(t.encode()).digest()
            out.append([((h[i] / 255.0) - 0.5) * 2 for i in range(8)])
        return out

    store = FileStore(_Client(), "real_test")
    store._collection = _Coll()
    configure(store, embed_fn)

    res = json.loads(ingest_file.invoke({"path": str(docx_path)}))
    check("返回 mode=inline", res["mode"] == "inline", str(res))
    check("summary 含文件名", "design.docx" in res["summary"])
    check("file_id 进缓存", res["file_id"] in _INLINE_CACHE)
    full = _INLINE_CACHE[res["file_id"]]
    ans = query_file.invoke({"file_id": res["file_id"], "question": "架构"})
    check("query 返回全文且含关键词", "RAG" in ans and ans.strip() == full.strip())

    print("\n" + "=" * 50)
    print(f"总计: {len(results)} 项, 通过 {sum(results)}, 失败 {len(results)-sum(results)}")
    sys.exit(0 if all(results) else 1)


if __name__ == "__main__":
    main()
